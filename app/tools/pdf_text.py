from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Protocol

from pydantic import BaseModel

from app.tools.text_cleaning import clean_text

MIN_TEXT_CHARACTERS = 40
EXPECTED_LABELS = (
    "legal name", "company name", "registration number", "tax id", "tax identification",
    "account number", "account no", "bank", "ifsc", "swift", "routing code", "currency",
    "total", "payment terms", "quotation", "invoice", "effective date", "address",
)


class TextQuality(BaseModel):
    char_count: int
    has_text: bool
    label_hits: int
    usable: bool
    reason: str


@dataclass
class ExtractedText:
    text: str
    pages: list[str]
    quality: TextQuality
    needs_ocr: bool = False


class OcrProvider(Protocol):
    def extract_text(self, source: Path | bytes) -> str: ...


class NullOcrProvider:
    def extract_text(self, source: Path | bytes) -> str:
        del source
        raise RuntimeError("OCR not configured")


def assess_text_quality(text: str) -> TextQuality:
    normalized = text.casefold()
    hits = sum(label in normalized for label in EXPECTED_LABELS)
    char_count = len(text.strip())
    has_text = char_count > 0
    usable = char_count >= MIN_TEXT_CHARACTERS and hits > 0
    reason = "usable text extracted" if usable else (
        "insufficient text" if char_count < MIN_TEXT_CHARACTERS else "no expected field labels found"
    )
    return TextQuality(
        char_count=char_count,
        has_text=has_text,
        label_hits=hits,
        usable=usable,
        reason=reason,
    )


def extract_text(source: Path | bytes) -> ExtractedText:
    suffix = source.suffix.lower() if isinstance(source, Path) else ""
    if suffix == ".pdf" or (isinstance(source, bytes) and source.startswith(b"%PDF")):
        pages = _extract_pdf_pages(source)
    elif suffix == ".docx" or (isinstance(source, bytes) and source.startswith(b"PK\x03\x04")):
        pages = [_extract_docx(source)]
    else:
        pages = [_extract_txt(source)]
    cleaned_pages = [clean_text(page) for page in pages]
    full_text = "\n\n".join(page for page in cleaned_pages if page)
    quality = assess_text_quality(full_text)
    return ExtractedText(
        text=full_text,
        pages=cleaned_pages,
        quality=quality,
        needs_ocr=not quality.usable,
    )


def extract_tables_text(source: Path | bytes) -> str:
    """Extract table cells from PDFs; callers should restrict use to poor bank/quote text."""
    try:
        import pdfplumber
    except ImportError as exc:
        raise RuntimeError("pdfplumber is required for PDF table extraction") from exc
    if isinstance(source, bytes):
        pdf = pdfplumber.open(BytesIO(source))
    else:
        pdf = pdfplumber.open(str(source))
    with pdf:
        rows: list[str] = []
        for page in pdf.pages:
            for table in page.extract_tables() or []:
                for row in table:
                    rows.append(" | ".join((cell or "").strip() for cell in row))
        return clean_text("\n".join(rows))


def _extract_pdf_pages(source: Path | bytes) -> list[str]:
    try:
        import fitz

        document = fitz.open(stream=source, filetype="pdf") if isinstance(source, bytes) else fitz.open(source)
        with document:
            pages = [page.get_text("text") for page in document]
        if any(page.strip() for page in pages):
            return pages
    except Exception:  # noqa: BLE001, S110 - malformed PDFs use the fallback parser
        pass
    try:
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(source)) if isinstance(source, bytes) else PdfReader(str(source))
        return [page.extract_text() or "" for page in reader.pages]
    except Exception as exc:
        raise ValueError("Unable to read PDF with PyMuPDF or pypdf") from exc


def _extract_docx(source: Path | bytes) -> str:
    try:
        from docx import Document

        document = Document(BytesIO(source)) if isinstance(source, bytes) else Document(str(source))
    except Exception as exc:
        raise ValueError("Unable to read DOCX document") from exc
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        paragraphs.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
    return "\n".join(paragraphs)


def _extract_txt(source: Path | bytes) -> str:
    data = source.read_bytes() if isinstance(source, Path) else source
    return data.decode("utf-8-sig", errors="replace")
