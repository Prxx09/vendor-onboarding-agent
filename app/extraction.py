"""Document text/OCR extraction and conservative vendor-field parsing."""
from __future__ import annotations

from io import BytesIO
from pathlib import Path
import re
from typing import Any


SUPPORTED_EXTENSIONS = {".pdf", ".png", ".jpg", ".jpeg", ".docx", ".txt"}
FIELD_LABELS = {
    "legal_name": "Legal name",
    "tax_id": "Tax ID / GSTIN",
    "pan": "PAN",
    "bank_account": "Bank account",
    "ifsc": "IFSC",
    "registered_address": "Registered address",
    "contact_email": "Contact email",
}


def _ocr_image(content: bytes) -> str:
    from PIL import Image

    with Image.open(BytesIO(content)) as image:
        if image.width * image.height > 25_000_000:
            raise ValueError("Image is too large for OCR (maximum 25 megapixels)")
    try:
        from rapidocr import RapidOCR
    except ImportError as exc:
        raise RuntimeError("OCR support is not installed. Run: pip install -r requirements.txt") from exc

    result = RapidOCR()(content)
    if result is None:
        return ""
    if hasattr(result, "txts"):
        return "\n".join(str(value) for value in (result.txts or []))
    rows = result[0] if isinstance(result, tuple) else result
    return "\n".join(str(row[1]) for row in (rows or []) if isinstance(row, (list, tuple)) and len(row) > 1)


def _pdf_text(content: bytes) -> tuple[str, str]:
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(content))
    if len(reader.pages) > 40:
        raise ValueError("PDF exceeds the 40-page extraction limit")
    text = "\n".join((page.extract_text() or "") for page in reader.pages).strip()
    if len(text) >= 40:
        return text, "pdf_text"

    # Scanned PDFs need rendering before OCR. PyMuPDF ships wheels on Windows,
    # so this does not require a separate Poppler or Tesseract installation.
    import fitz

    document = fitz.open(stream=content, filetype="pdf")
    pages = []
    for page in document:
        pixmap = page.get_pixmap(matrix=fitz.Matrix(1.8, 1.8), alpha=False)
        pages.append(_ocr_image(pixmap.tobytes("png")))
    return "\n".join(pages).strip(), "ocr"


def _docx_text(content: bytes) -> str:
    from docx import Document

    document = Document(BytesIO(content))
    values = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        values.extend(cell.text for row in table.rows for cell in row.cells)
    return "\n".join(value for value in values if value.strip())


def extract_text(name: str, content: bytes) -> tuple[str, str]:
    extension = Path(name).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported document type: {extension or 'unknown'}")
    if extension == ".pdf":
        return _pdf_text(content)
    if extension in {".png", ".jpg", ".jpeg"}:
        return _ocr_image(content), "ocr"
    if extension == ".docx":
        return _docx_text(content), "docx_text"
    return content.decode("utf-8", errors="replace"), "plain_text"


def _label_value(text: str, labels: list[str]) -> str:
    joined = "|".join(re.escape(label) for label in labels)
    match = re.search(rf"(?im)^\s*(?:{joined})\s*[:\-]\s*(.+?)\s*$", text)
    return match.group(1).strip() if match else ""


def _first(pattern: str, text: str) -> str:
    match = re.search(pattern, text, re.IGNORECASE)
    return match.group(0).strip() if match else ""


def parse_vendor_fields(text: str) -> tuple[dict[str, str], dict[str, int]]:
    fields = {
        "legal_name": _label_value(text, ["Legal name", "Vendor name", "Supplier name", "Company name"]),
        "tax_id": _label_value(text, ["Tax ID / GSTIN", "Tax ID", "GSTIN", "GST No"]),
        "pan": _label_value(text, ["PAN", "PAN No", "Permanent Account Number"]),
        "bank_account": _label_value(text, ["Bank account", "Account number", "A/C No", "Account"]),
        "ifsc": _label_value(text, ["IFSC", "IFSC code"]),
        "registered_address": _label_value(text, ["Registered address", "Business address", "Address"]),
        "contact_email": _label_value(text, ["Contact email", "Email", "E-mail"]),
    }
    fields["tax_id"] = fields["tax_id"] or _first(r"\b\d{2}[A-Z]{5}\d{4}[A-Z][A-Z0-9]Z[A-Z0-9]\b", text)
    fields["pan"] = fields["pan"] or _first(r"\b[A-Z]{5}\d{4}[A-Z]\b", text)
    fields["ifsc"] = fields["ifsc"] or _first(r"\b[A-Z]{4}0[A-Z0-9]{6}\b", text)
    fields["contact_email"] = fields["contact_email"] or _first(r"[A-Z0-9._%+\-]+@[A-Z0-9.\-]+\.[A-Z]{2,}", text)
    fields = {key: value.strip(" .") for key, value in fields.items() if value.strip()}
    confidence = {key: (96 if key in {"tax_id", "pan", "ifsc", "contact_email"} else 90) for key in fields}
    return fields, confidence


def document_types(name: str, text: str) -> list[str]:
    sample = f"{name} {text[:1000]}".lower()
    rules = [
        ("business_registration", ("registration", "incorporation", "udyam", "msme")),
        ("tax_certificate", ("gstin", "tax certificate", "tax id", "pan")),
        ("bank_proof", ("cancelled cheque", "bank statement", "ifsc", "bank account")),
        ("address_proof", ("address proof", "registered address", "utility bill")),
        ("compliance_declaration", ("compliance", "declaration", "questionnaire")),
    ]
    found = [kind for kind, terms in rules if any(term in sample for term in terms)]
    return found or ["supporting_document"]


def merge_extractions(items: list[dict[str, Any]]) -> dict[str, Any]:
    fields: dict[str, str] = {}
    confidence: dict[str, int] = {}
    sources: dict[str, str] = {}
    for item in items:
        for key, value in item["fields"].items():
            score = item["field_confidence"].get(key, 0)
            if not fields.get(key) or score > confidence.get(key, 0):
                fields[key] = value
                confidence[key] = score
                sources[key] = item["name"]
    expected = len(FIELD_LABELS)
    extraction_confidence = round(sum(confidence.values()) / expected) if expected else 0
    return {
        "extracted_fields": fields,
        "field_confidence": confidence,
        "field_sources": sources,
        "extraction_confidence": extraction_confidence,
        "missing_fields": [key for key in FIELD_LABELS if not fields.get(key)],
    }


def extract_documents(documents: list[tuple[str, bytes]]) -> dict[str, Any]:
    results = []
    for name, content in documents:
        try:
            text, method = extract_text(name, content)
        except Exception as exc:
            raise ValueError(f"Could not extract {name}: {exc}") from exc
        fields, confidence = parse_vendor_fields(text)
        detected_types = document_types(name, text)
        results.append({
            "name": name,
            "method": method,
            "document_type": detected_types[0],
            "detected_types": detected_types,
            "fields": fields,
            "field_confidence": confidence,
            "characters_extracted": len(text),
        })
    return {**merge_extractions(results), "documents": results}
