import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from pydantic import ValidationError

from app.agents.document_understanding import DocumentInput, DocumentUnderstandingAgent
from app.config import Settings
from app.llm.errors import LLMProviderError
from app.llm.groq_provider import GroqProvider
from app.llm.mock_provider import MockProvider
from app.repositories.memory import InMemoryRepository
from app.repositories.storage import LocalDocumentStorage
from app.rules.confidence import evaluate_extracted_field
from app.rules.document_fields import DocumentType
from app.schemas.common import AgentStatus
from app.schemas.domain import Document, RequestType, VendorRequest
from app.schemas.extraction import ExtractedFieldStatus, LLMExtractionOutput
from app.services.extraction_service import ExtractionService
from app.tools.pdf_text import NullOcrProvider, extract_text
from app.tools.text_cleaning import clean_text

FIXTURE_DIR = Path(__file__).parent / "fixtures"


def canned(name: str) -> dict:
    return json.loads((FIXTURE_DIR / "llm" / f"{name}.json").read_text(encoding="utf-8"))


def run_fixture(name: str, doc_type: DocumentType | None = None, response_name: str | None = None):
    text = (FIXTURE_DIR / "documents" / f"{name}.txt").read_text(encoding="utf-8")
    provider = MockProvider({"document_extraction": canned(response_name or name)})
    result = DocumentUnderstandingAgent(provider).run(
        DocumentInput(
            document_id=f"doc-{name}",
            file_name=f"{name}.txt",
            text=text,
            declared_type=doc_type,
        )
    )
    return result, provider, text


@pytest.mark.parametrize(
    ("name", "doc_type"),
    [
        ("business_registration", DocumentType.BUSINESS_REGISTRATION),
        ("tax_certificate", DocumentType.TAX_CERTIFICATE),
        ("bank_document", DocumentType.BANK_DOCUMENT),
        ("quotation", DocumentType.QUOTATION),
    ],
)
def test_valid_document_type_extraction(name, doc_type):
    result, _, _ = run_fixture(name, doc_type)
    assert result.data is not None
    assert result.data.document_type == doc_type
    assert result.data.fields
    if doc_type == DocumentType.QUOTATION:
        assert result.data.line_items[0].description.value == "Demo paper pack"


def test_missing_required_field_is_flagged():
    result, _, _ = run_fixture("missing_field", DocumentType.TAX_CERTIFICATE)
    assert result.data.missing_required_fields == ["tax_id"]
    assert result.status == AgentStatus.NEEDS_MANUAL_REVIEW


def test_hallucinated_value_is_unverified():
    source = "SYNTHETIC TAX CERTIFICATE\nLegal Name: Demo Company Incorporated\nTax ID: DEMO-100"
    field = evaluate_extracted_field(
        document_id="d1",
        name="legal_name",
        value="Invented Company",
        evidence_quote="Legal Name: Invented Company",
        source_text=source,
    )
    assert field.status == ExtractedFieldStatus.UNVERIFIED
    assert field.confidence == 0
    assert field.needs_review


def test_quote_not_found_gets_reduced_confidence():
    field = evaluate_extracted_field(
        document_id="d1",
        name="legal_name",
        value="Demo Company",
        evidence_quote="This quote is not in the document",
        source_text="Legal Name: Demo Company",
    )
    assert field.status == ExtractedFieldStatus.FOUND
    assert field.confidence == 0.7
    assert field.evidence.quote == "Legal Name: Demo Company"
    needs_review = evaluate_extracted_field(
        document_id="d1",
        name="legal_name",
        value="Demo Company",
        evidence_quote="bad quote",
        source_text="Legal Name: Demo Company",
        threshold=0.8,
    )
    assert needs_review.needs_review


def test_confidence_missing_and_exact_evidence_branches():
    missing = evaluate_extracted_field(
        document_id="d1", name="tax_id", value=None, evidence_quote=None, source_text="Tax ID: x"
    )
    exact = evaluate_extracted_field(
        document_id="d1",
        name="effective_date",
        value="15 January 2024",
        evidence_quote="Effective Date: 2024-01-15",
        source_text="Effective Date: 2024-01-15; total USD 1,234.50",
    )
    numeric = evaluate_extracted_field(
        document_id="d1",
        name="total",
        value="1234.50",
        evidence_quote="total USD 1,234.50",
        source_text="Effective Date: 2024-01-15; total USD 1,234.50",
    )
    assert missing.status == ExtractedFieldStatus.MISSING and missing.confidence == 0
    assert exact.status == ExtractedFieldStatus.FOUND and exact.confidence == 1
    assert numeric.status == ExtractedFieldStatus.FOUND and numeric.confidence == 1


def test_text_cleaning_normalizes_and_truncates():
    assert clean_text("  alpha\x00\r\n  beta  ") == "alpha\nbeta"
    assert clean_text("abcdefghijklmnopqrstuvwxyz", max_length=20).endswith("[TEXT TRUNCATED]")


def test_prompt_injection_is_data_and_cannot_set_risk_or_status(caplog):
    result, provider, text = run_fixture("prompt_injection", DocumentType.TAX_CERTIFICATE)
    call = provider.calls[0]
    assert f"<DOCUMENT_DATA>\n{clean_text(text)}\n</DOCUMENT_DATA>" in call["user"]
    assert "untrusted DATA" in call["user"]
    assert result.data is not None
    assert not hasattr(result.data, "status")
    assert not hasattr(result.data, "risk_score")
    assert "Ignore previous instructions" not in caplog.text


def test_bank_account_is_masked_in_prompt_output_and_logs(caplog):
    result, provider, text = run_fixture("bank_document", DocumentType.BANK_DOCUMENT)
    prompt = provider.calls[0]["user"]
    assert "123456789012" not in prompt
    assert "****9012" in prompt
    field = next(field for field in result.data.fields if field.name == "account_number")
    assert field.value == "****9012"
    assert field.status == ExtractedFieldStatus.UNVERIFIED
    assert "123456789012" not in caplog.text
    assert text  # source is used locally for evidence checks, not logged


def test_declared_type_wins_and_model_mismatch_is_uncertainty():
    result, _, _ = run_fixture(
        "tax_certificate", DocumentType.BUSINESS_REGISTRATION, response_name="tax_certificate"
    )
    assert result.data.document_type == DocumentType.BUSINESS_REGISTRATION
    assert any("differed from declared type" in item for item in result.data.uncertainties)


def test_provider_failure_returns_manual_review_without_fabrication():
    provider = MockProvider(fail_with=LLMProviderError("provider unavailable"))
    result = DocumentUnderstandingAgent(provider).run(
        {"document_id": "d1", "file_name": "x.txt", "text": "A short source"}
    )
    assert result.status == AgentStatus.NEEDS_MANUAL_REVIEW
    assert result.data is None


class FakeCompletions:
    def __init__(self, contents):
        self.contents = iter(contents)

    def create(self, **kwargs):
        content = next(self.contents)
        if isinstance(content, Exception):
            raise content
        return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=content))])


class FakeClient:
    def __init__(self, contents):
        self.chat = SimpleNamespace(completions=FakeCompletions(contents))


def test_malformed_json_retries_then_fails_as_manual_review():
    settings = Settings(groq_api_key="fake", groq_max_retries=0, _env_file=None)
    provider = GroqProvider(settings, FakeClient(["not json", "still not json"]))
    agent = DocumentUnderstandingAgent(provider)
    result = agent.run(
        {"document_id": "d1", "file_name": "x.txt", "text": "Legal Name: Synthetic Example Vendor"}
    )
    assert result.status == AgentStatus.NEEDS_MANUAL_REVIEW
    assert result.data is None


def test_groq_provider_network_failure_returns_manual_review():
    settings = Settings(groq_api_key="fake", groq_max_retries=0, _env_file=None)
    provider = GroqProvider(settings, FakeClient([RuntimeError("offline")]))
    result = DocumentUnderstandingAgent(provider).run(
        {"document_id": "d2", "file_name": "x.txt", "text": "Legal Name: Synthetic Example Vendor"}
    )
    assert result.status == AgentStatus.NEEDS_MANUAL_REVIEW
    assert result.data is None


def test_pdf_extraction_with_pymupdf():
    fitz = pytest.importorskip("fitz")
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Legal Name: Example Vendor\nTax ID: DEMO-9912\n" + "sample " * 12)
    extracted = extract_text(pdf.tobytes())
    assert "Example Vendor" in extracted.text
    assert extracted.quality.usable
    assert len(extracted.pages) == 1
    pdf.close()


def test_docx_extraction_from_bytes():
    from io import BytesIO

    from docx import Document as DocxDocument

    document = DocxDocument()
    document.add_paragraph("Legal Name: Example Vendor with enough text for the quality check")
    stream = BytesIO()
    document.save(stream)
    extracted = extract_text(stream.getvalue())
    assert "Example Vendor" in extracted.text
    assert extracted.quality.usable


def test_unusable_text_sets_needs_ocr_and_null_provider_explains():
    extracted = extract_text(b"tiny")
    assert extracted.needs_ocr
    assert not extracted.quality.usable
    with pytest.raises(RuntimeError, match="OCR not configured"):
        NullOcrProvider().extract_text(b"tiny")


def test_extraction_service_persists_fields_run_and_audit():
    text = (FIXTURE_DIR / "documents" / "tax_certificate.txt").read_text(encoding="utf-8")
    repo = InMemoryRepository()
    request = repo.create_request(VendorRequest(request_type=RequestType.NEW_SUPPLIER))
    storage = LocalDocumentStorage(Path(__file__).parent)
    storage_path = f"extraction-{request.id}.txt"
    try:
        storage.upload(storage_path, text.encode("utf-8"), "text/plain")
        document = repo.add_document(
            Document(
                request_id=request.id,
                doc_type="TAX_CERTIFICATE",
                file_name="tax_certificate.txt",
                storage_path=storage_path,
            )
        )
        provider = MockProvider({"document_extraction": canned("tax_certificate")})
        service = ExtractionService(repo, DocumentUnderstandingAgent(provider), storage)
        result = service.extract_document(document.id)
        assert result.status == AgentStatus.OK
        assert len(repo.list_extracted_fields(request.id)) == 5
        assert len(repo.agent_runs) == 1
        assert any(
            event.event_type == "DOCUMENT_EXTRACTION_COMPLETED"
            for event in repo.list_audit_events(request.id)
        )
        cached_document = repo.get_document(document.id)
        assert cached_document.text_content == clean_text(text)
        assert cached_document.text_quality["usable"]
    finally:
        (Path(__file__).parent / storage_path).unlink(missing_ok=True)


def test_llm_output_schema_forbids_decision_fields():
    with pytest.raises(ValidationError):
        LLMExtractionOutput.model_validate(
            {"document_type": "UNKNOWN", "risk_score": 0, "status": "APPROVED"}
        )
